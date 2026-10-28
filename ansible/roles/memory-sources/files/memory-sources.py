#!/usr/bin/env python3
"""Write Markdown copies of workspace files that OpenClaw's memory search cannot read.

OpenClaw's built-in memory search indexes Markdown only. This keeps the rest of what
qmd used to search findable: plain text and CSV under notes/, and the text inside
PDFs, images, Word and Excel files anywhere outside dot folders.

Runs in a throwaway container (see the memory-sources role): the agent's workspace is
mounted read-only at /source, its copy folder writable at /out. Each copy is
/out/<workspace path>.md and starts with the workspace path of its original.
Unchanged sources are skipped; copies whose source is gone or no longer readable
are removed, so a copy is never stale. A file that fails to read is given up on
(until it changes) after one failure, or after three when the cause may pass (a
timeout, a full disk) or the run was stopped while reading it; giving up makes the
run exit 3 once, and the unit's failure alert tells the owner.
"""

import errno
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

SOURCE = Path('/source')
OUT = Path('/out')
MANIFEST = OUT / '.manifest.json'
# Bump to rewrite every copy after changing what or how this script extracts.
VERSION = 1
TIMEOUT = 60
# Tries before giving up on a file whose failure may pass (timeout, full disk, stopped run).
ATTEMPTS = 3
OCR_LANGUAGES = 'eng+deu'

TEXT = {'.txt', '.csv'}
PDF = {'.pdf'}
IMAGE = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.bmp', '.tiff', '.tif'}
DOCX = {'.docx'}
XLSX = {'.xlsx'}


def kind(relative):
    suffix = relative.suffix.lower()
    if suffix in TEXT:
        # qmd indexed notes/**/*.{txt,csv}; text elsewhere stays unindexed as before.
        return 'text' if relative.parts[0] == 'notes' else None
    for name, suffixes in (('pdf', PDF), ('image', IMAGE), ('docx', DOCX), ('xlsx', XLSX)):
        if suffix in suffixes:
            return name
    return None


def sources():
    """Workspace files to copy, skipping dot folders and symbolic links."""
    found = {}
    for root, dirs, files in os.walk(SOURCE):
        dirs[:] = sorted(d for d in dirs if not d.startswith('.'))
        for name in files:
            path = Path(root) / name
            if path.is_symlink() or not path.is_file():
                continue
            relative = path.relative_to(SOURCE)
            file_kind = kind(relative)
            if file_kind:
                found[relative.as_posix()] = (path, file_kind)
    return found


def run(args):
    result = subprocess.run(args, capture_output=True, text=True, timeout=TIMEOUT)
    if result.returncode != 0:
        raise RuntimeError(f'{args[0]} exited {result.returncode}: {result.stderr.strip()[:300]}')
    return result.stdout.strip()


def ocr(path):
    return run(['tesseract', str(path), 'stdout', '-l', OCR_LANGUAGES])


def pdf_text(path):
    text = run(['pdftotext', str(path), '-'])
    words = text.split()
    # Scanned PDFs have no text layer: OCR the first ten pages instead.
    if len(words) < 5 or sum(c.isalpha() for c in text) / len(text) < 0.5:
        with tempfile.TemporaryDirectory() as pages:
            run(['pdftoppm', '-r', '200', '-l', '10', str(path), os.path.join(pages, 'page')])
            scanned = '\n\n'.join(ocr(page) for page in sorted(Path(pages).glob('page-*.ppm')))
        if len(scanned.split()) > len(words):
            return scanned
    return text


def docx_text(path):
    from docx import Document
    return '\n'.join(p.text for p in Document(str(path)).paragraphs if p.text.strip())


def xlsx_text(path):
    from openpyxl import load_workbook
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f'[{sheet.title}]')
        for row in sheet.iter_rows(values_only=True):
            cells = ['' if cell is None else str(cell) for cell in row]
            if any(cells):
                lines.append('\t'.join(cells))
    workbook.close()
    return '\n'.join(lines)


EXTRACT = {
    'text': lambda path: path.read_bytes().decode('utf-8', errors='replace'),
    'pdf': pdf_text,
    'image': ocr,
    'docx': docx_text,
    'xlsx': xlsx_text,
}
DESCRIPTION = {
    'text': 'Copy of the workspace file',
    'pdf': 'Text extracted from the PDF',
    'image': 'Text read (OCR) from the image',
    'docx': 'Text extracted from the Word document',
    'xlsx': 'Cells extracted from the Excel workbook',
}


def write(target, content):
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUT / '.writing.tmp'  # short fixed name: long file names still fit
    temporary.write_text(content, errors='replace')
    temporary.replace(target)


def shown(relative):
    """The path as text, even when the file name is not valid UTF-8."""
    return relative.encode('utf-8', 'surrogateescape').decode('utf-8', 'replace')


def main():
    for tool in ('pdftotext', 'pdftoppm', 'tesseract'):
        if not shutil.which(tool):
            raise SystemExit(f'ERROR: {tool} is missing from the image')
    try:
        manifest = json.loads(MANIFEST.read_text())
    except FileNotFoundError:
        manifest = {}
    if manifest.get('version') != VERSION:
        manifest = {'version': VERSION, 'files': {}, 'unreadable': {}}
    known, unreadable = manifest['files'], manifest['unreadable']

    def save():
        write(MANIFEST, json.dumps(manifest))

    # The previous run stopped while reading this file: out of memory or time on that
    # file, or just a reboot. Read it again, and give up after the third stop.
    interrupted = manifest.pop('reading', None)
    newly_unreadable = 0

    def give_up(relative, signature, error, attempts=1, transient=False):
        """Record a failure; True once the file is skipped until it changes."""
        nonlocal newly_unreadable
        final = not transient or attempts >= ATTEMPTS
        unreadable[relative] = {'signature': signature, 'error': error, 'attempts': attempts, 'final': final}
        known.pop(relative, None)
        if final:
            newly_unreadable += 1

    if interrupted and interrupted.get('attempts', 1) >= ATTEMPTS:
        give_up(interrupted['path'], interrupted['signature'],
                'runs stopped while reading this file (out of memory or time)', interrupted['attempts'])

    found = sources()
    written = 0
    for relative, (path, file_kind) in sorted(found.items()):
        try:
            stat = path.stat()
        except FileNotFoundError:
            continue  # deleted during the walk; the next run removes its copy
        signature = [stat.st_mtime_ns, stat.st_size]
        failure = unreadable.get(relative, {})
        if known.get(relative) == signature or (failure.get('signature') == signature and failure.get('final', True)):
            continue
        attempts = 1
        if interrupted and interrupted['path'] == relative and interrupted['signature'] == signature:
            attempts = interrupted.get('attempts', 1) + 1
        elif failure.get('signature') == signature:
            attempts = failure.get('attempts', 1) + 1
        manifest['reading'] = {'path': relative, 'signature': signature, 'attempts': attempts}
        save()
        try:
            text = EXTRACT[file_kind](path)
            write(OUT / f'{relative}.md',
                  f'# {shown(relative)}\n\n{DESCRIPTION[file_kind]} `{shown(relative)}` in this agent\'s workspace.\n\n{text}\n')
        except Exception as error:  # one unreadable file must not stop the others
            if path.exists():
                transient = isinstance(error, subprocess.TimeoutExpired) or (
                    isinstance(error, OSError) and error.errno in (errno.ENOSPC, errno.EIO))
                give_up(relative, signature, f'{type(error).__name__}: {error}', attempts, transient)
            else:
                known.pop(relative, None)
        else:
            unreadable.pop(relative, None)
            known[relative] = signature
            written += 1
        manifest.pop('reading')
        save()  # progress survives a crash or a time limit

    removed = 0
    for relative in sorted(set(known) - set(found)):
        known.pop(relative)
    for relative in sorted(set(unreadable) - set(found)):
        unreadable.pop(relative)
    for copy in sorted(OUT.rglob('*.md')):
        relative = copy.relative_to(OUT).as_posix()[:-3]
        if relative not in known:
            copy.unlink()
            removed += 1
    for directory in sorted((d for d in OUT.rglob('*') if d.is_dir()), reverse=True):
        if not any(directory.iterdir()):
            directory.rmdir()

    save()
    print(f'{len(known)} copies: {written} written, {removed} removed, {len(unreadable)} unreadable')
    for relative, entry in sorted(unreadable.items()):
        label = 'UNREADABLE' if entry.get('final', True) else f'RETRY ({entry["attempts"]}/{ATTEMPTS})'
        print(f'{label} {shown(relative)}: {entry["error"]}')
    return 3 if newly_unreadable else 0


if __name__ == '__main__':
    sys.exit(main())
